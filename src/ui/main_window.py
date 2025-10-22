"""
Main window for NeuronDB application
"""

import tkinter as tk
from tkinter import ttk, messagebox, filedialog
import customtkinter as ctk
from typing import Dict, Any, Optional
import sys
import os
from pathlib import Path
import pandas as pd
from datetime import datetime
from PIL import Image, ImageTk
import threading

# Add parent directory to path for imports
sys.path.append(str(Path(__file__).parent.parent))

from database.connection import DatabaseConnection, ConnectionManager
from ai.assistant import NeuronDBAI
from utils.helpers import setup_logging
from utils.theme_manager import theme_manager
from utils.config_manager import config_manager, apply_startup_theme
from ui.connection_dialog import ConnectionDialog
from ui.query_panel import QueryPanel
from ui.schema_browser import SchemaBrowser
from ui.psql_terminal import PSQLTerminal
from ui.db_diagram_view import DBDiagramView
from ui.config_view import ConfigView
from config import Config

# Set appearance mode and color theme
ctk.set_appearance_mode("light")  # Modes: "System" (standard), "Dark", "Light"
ctk.set_default_color_theme("blue")  # Themes: "blue" (standard), "green", "dark-blue"

# Configure CTk default colors using theme manager
def apply_theme_to_ctk():
    """Apply current theme colors to CustomTkinter defaults"""
    try:
        # Apply background colors
        ctk.ThemeManager.theme["CTkFrame"]["fg_color"] = [
            theme_manager.get_color("background.main"), 
            theme_manager.get_color("background.main")
        ]
        ctk.ThemeManager.theme["CTk"]["fg_color"] = [
            theme_manager.get_color("background.main"), 
            theme_manager.get_color("background.main")
        ]
        
        # Apply button colors
        ctk.ThemeManager.theme["CTkButton"]["fg_color"] = [
            theme_manager.get_color("buttons.primary_bg"), 
            theme_manager.get_color("buttons.primary_bg")
        ]
        ctk.ThemeManager.theme["CTkButton"]["hover_color"] = [
            theme_manager.get_color("buttons.primary_hover"), 
            theme_manager.get_color("buttons.primary_hover")
        ]
        ctk.ThemeManager.theme["CTkButton"]["text_color"] = [
            theme_manager.get_color("buttons.primary_text"), 
            theme_manager.get_color("buttons.primary_text")
        ]
        
        # Apply scrollable frame colors
        ctk.ThemeManager.theme["CTkScrollableFrame"]["fg_color"] = [
            theme_manager.get_color("background.main"), 
            theme_manager.get_color("background.main")
        ]
        
        # Apply entry/input colors
        ctk.ThemeManager.theme["CTkEntry"]["fg_color"] = [
            theme_manager.get_color("editor.background"), 
            theme_manager.get_color("editor.background")
        ]
        ctk.ThemeManager.theme["CTkEntry"]["border_color"] = [
            theme_manager.get_color("accent.main"), 
            theme_manager.get_color("accent.main")
        ]
        ctk.ThemeManager.theme["CTkEntry"]["text_color"] = [
            theme_manager.get_color("text.primary"), 
            theme_manager.get_color("text.primary")
        ]
        
        # Apply label colors
        ctk.ThemeManager.theme["CTkLabel"]["text_color"] = [
            theme_manager.get_color("text.primary"), 
            theme_manager.get_color("text.primary")
        ]
        
        # Apply optionmenu colors
        ctk.ThemeManager.theme["CTkOptionMenu"]["fg_color"] = [
            theme_manager.get_color("buttons.secondary_bg"), 
            theme_manager.get_color("buttons.secondary_bg")
        ]
        ctk.ThemeManager.theme["CTkOptionMenu"]["button_color"] = [
            theme_manager.get_color("buttons.primary_bg"), 
            theme_manager.get_color("buttons.primary_bg")
        ]
        ctk.ThemeManager.theme["CTkOptionMenu"]["button_hover_color"] = [
            theme_manager.get_color("buttons.primary_hover"), 
            theme_manager.get_color("buttons.primary_hover")
        ]
        
        print("Applied theme to CTk components")
    except Exception as e:
        print(f"Error applying theme to CTk: {e}")

# Apply theme
apply_theme_to_ctk()

class NeuronDBApp(ctk.CTk):
    """Main application window for NeuronDB"""
    
    def __init__(self):
        super().__init__()
        
        # Set up logging
        self.logger = setup_logging()
        
        # Initialize core components
        self.db_connection = DatabaseConnection()

        self.connection_manager = ConnectionManager()
        self.ai_assistant = None
        self.current_schema = {}
        
        # Configure main window
        self.title("NeuronDB")
        self.geometry("1600x1100")
        self.minsize(1400, 950)
        
        # Initialize theme system with user's preferred theme
        print(f"Available themes: {theme_manager.list_available_themes()}")
        preferred_theme = config_manager.config.selected_theme
        theme_manager.initialize_with_fallback(preferred_theme)
        print(f"Current theme: {theme_manager.get_theme_name()}")
        
        # Apply theme before creating UI
        self.apply_theme()
        
        # Configure window for better text display
        self.tk.call('tk', 'scaling', 1.0)  # Ensure consistent scaling
        
        # Set window icon (desktop icon)
        try:
            if Config.DESKTOP_ICON.exists():
                icon_image = Image.open(Config.DESKTOP_ICON)
                icon_photo = ImageTk.PhotoImage(icon_image)
                self.iconphoto(True, icon_photo)
                # Keep reference to prevent garbage collection
                self._icon_photo = icon_photo
        except Exception as e:
            self.logger.warning(f"Could not load desktop icon: {e}")
        
        # Configure grid weights
        self.grid_columnconfigure(0, weight=1)
        self.grid_rowconfigure(0, weight=1)
        self.grid_rowconfigure(1, weight=0)
        
        # Create UI components
        self.create_menu_bar()
        self.create_main_interface()
        self.create_status_bar()
        
        # Initialize AI assistant
        self.init_ai_assistant()
        
        # Bind close event
        self.protocol("WM_DELETE_WINDOW", self.on_closing)
        
        self.logger.info("NeuronDB application initialized")
    
    def apply_theme(self):
        """Apply current theme to all components"""
        # Re-apply theme to CustomTkinter defaults
        apply_theme_to_ctk()
        
        # Set window background
        self.configure(fg_color=theme_manager.get_color("background.main"))
        
        # Update main interface components if they exist
        if hasattr(self, 'left_frame'):
            self.left_frame.configure(fg_color=theme_manager.get_color("sidebar.background"))
        
        if hasattr(self, 'main_tabs'):
            self.main_tabs.configure(
                fg_color=theme_manager.get_color("background.main"),
                segmented_button_fg_color=theme_manager.get_color("sidebar.background"),
                segmented_button_selected_color=theme_manager.get_color("buttons.primary_bg"),
                segmented_button_selected_hover_color=theme_manager.get_color("buttons.primary_hover"),
                segmented_button_unselected_color=theme_manager.get_color("sidebar.background"),
                segmented_button_unselected_hover_color=theme_manager.get_color("sidebar.header"),
                text_color=theme_manager.get_color("text.primary"),
                text_color_disabled=theme_manager.get_color("text.secondary")
            )
        
        if hasattr(self, 'query_notebook'):
            self.query_notebook.configure(
                fg_color=theme_manager.get_color("background.main"),
                segmented_button_fg_color=theme_manager.get_color("sidebar.background")
            )
        
        if hasattr(self, 'status_frame'):
            self.status_frame.configure(fg_color=theme_manager.get_color("card"))
        
        # Update schema browser theme
        if hasattr(self, 'schema_browser'):
            self.schema_browser.apply_theme()
        
        # Update query panel theme  
        if hasattr(self, 'query_panel'):
            self.query_panel.apply_theme()
            
        # Update config view theme
        if hasattr(self, 'config_view'):
            self.config_view.apply_theme()
        
        # Force UI refresh
        self.update_idletasks()
    
    def switch_theme(self, theme_name: str):
        """Switch to a different theme and update all components"""
        from utils.theme_manager import theme_manager
        from utils.config_manager import config_manager
        
        try:
            if theme_manager.set_theme(theme_name):
                print(f"Successfully loaded theme: {theme_name}")
                
                # Save theme preference to config
                config_manager.config.selected_theme = theme_name
                config_manager.save_config()
                
                # Apply theme to all components
                self.apply_theme()
                
                # Also update the config view if it exists
                if hasattr(self, 'config_view'):
                    self.config_view.refresh_theme_selector()
                
                messagebox.showinfo("Theme Changed", f"Theme switched to: {theme_manager.get_theme_name()}")
                print(f"Theme applied and saved as default")
            else:
                messagebox.showerror("Error", f"Could not load theme: {theme_name}")
        except Exception as e:
            print(f"Error switching theme: {e}")
            messagebox.showerror("Error", f"Error switching theme: {str(e)}")
    
    def create_menu_bar(self):
        """Create the application menu bar"""
        menubar = tk.Menu(self)
        self.config(menu=menubar)
        
        # File menu
        file_menu = tk.Menu(menubar, tearoff=0)
        menubar.add_cascade(label="File", menu=file_menu)
        file_menu.add_command(label="New Connection...", command=self.show_connection_dialog)
        file_menu.add_separator()
        file_menu.add_command(label="Exit", command=self.on_closing)
        
        # Connection menu
        connection_menu = tk.Menu(menubar, tearoff=0)
        menubar.add_cascade(label="Connection", menu=connection_menu)
        connection_menu.add_command(label="Connect...", command=self.show_connection_dialog)
        connection_menu.add_command(label="Disconnect", command=self.disconnect_database)
        connection_menu.add_separator()
        connection_menu.add_command(label="Refresh Schema", command=self.refresh_schema)
        
        # Query menu
        query_menu = tk.Menu(menubar, tearoff=0)
        menubar.add_cascade(label="Query", menu=query_menu)
        query_menu.add_command(label="Execute Query", command=self.execute_current_query)
        query_menu.add_command(label="Clear Query", command=self.clear_query)
        query_menu.add_separator()
        query_menu.add_command(label="Format Query", command=self.format_current_query)
        
        # AI menu
        ai_menu = tk.Menu(menubar, tearoff=0)
        menubar.add_cascade(label="AI", menu=ai_menu)
        ai_menu.add_command(label="Generate Query...", command=self.show_ai_dialog)
        ai_menu.add_command(label="Explain Query", command=self.explain_current_query)
        ai_menu.add_separator()
        ai_menu.add_command(label="Clear Chat History", command=self.clear_ai_history)
        
        # Theme menu
        theme_menu = tk.Menu(menubar, tearoff=0)
        menubar.add_cascade(label="Theme", menu=theme_menu)
        
        # Add available themes to menu
        from utils.theme_manager import theme_manager
        available_themes = theme_manager.list_available_themes()
        theme_display_names = theme_manager.get_theme_display_names()
        
        for theme_file in available_themes:
            display_name = theme_display_names.get(theme_file, theme_file.title())
            theme_menu.add_command(
                label=display_name, 
                command=lambda t=theme_file: self.switch_theme(t)
            )
        
        # Help menu
        help_menu = tk.Menu(menubar, tearoff=0)
        menubar.add_cascade(label="Help", menu=help_menu)
        help_menu.add_command(label="About", command=self.show_about)
    
    def create_main_interface(self):
        """Create the main interface with tabbed layout"""
        # Create main tabbed interface
        self.main_tabs = ctk.CTkTabview(
            self, 
            fg_color=theme_manager.get_color("background.main"),
            segmented_button_fg_color=theme_manager.get_color("sidebar.background"),
            segmented_button_selected_color=theme_manager.get_color("buttons.primary_bg"),
            segmented_button_selected_hover_color=theme_manager.get_color("buttons.primary_hover"),
            segmented_button_unselected_color=theme_manager.get_color("sidebar.background"),
            segmented_button_unselected_hover_color=theme_manager.get_color("sidebar.header"),
            text_color=theme_manager.get_color("text.primary"),
            text_color_disabled=theme_manager.get_color("text.secondary")
        )
        self.main_tabs.grid(row=0, column=0, sticky="nsew", padx=10, pady=(10, 0))
        
        # Add tabs
        self.main_tabs.add("DB Query")
        self.main_tabs.add("DB View")
        self.main_tabs.add("Config")
        
        # Create DB Query tab (current app layout)
        self.create_db_query_tab()
        
        # Create DB View tab (DBML diagram)
        self.db_diagram_view = DBDiagramView(self.main_tabs.tab("DB View"))
        self.db_diagram_view.pack(fill="both", expand=True)
        
        # Create Config tab
        self.config_view = ConfigView(self.main_tabs.tab("Config"), main_window=self)
        self.config_view.pack(fill="both", expand=True)
    
    def create_db_query_tab(self):
        """Create the DB Query tab with the main query interface"""
        db_query_tab = self.main_tabs.tab("DB Query")
        
        # Configure grid for DB Query tab
        db_query_tab.grid_columnconfigure(0, weight=1)
        db_query_tab.grid_rowconfigure(0, weight=1)
        
        # Create main paned window
        self.main_paned = tk.PanedWindow(db_query_tab, orient=tk.HORIZONTAL, sashrelief=tk.RAISED, sashwidth=6)
        self.main_paned.grid(row=0, column=0, sticky="nsew", padx=0, pady=0)
        
        # Left panel - Schema Browser
        self.left_frame = ctk.CTkFrame(self.main_paned, width=350, fg_color=theme_manager.get_color("sidebar.background"), corner_radius=8)
        self.main_paned.add(self.left_frame, width=350, minsize=280)
        
        # Schema browser (will be fully initialized after query panel)
        self.schema_browser = None
        
        # Right panel - Query and results
        self.right_paned = tk.PanedWindow(self.main_paned, orient=tk.VERTICAL, sashrelief=tk.RAISED, sashwidth=6)
        self.main_paned.add(self.right_paned, minsize=500)
        
        # Query panel
        self.query_frame = ctk.CTkFrame(self.right_paned, height=400, corner_radius=8)
        self.right_paned.add(self.query_frame, height=400, minsize=250)
        
        # Create tabbed interface for query tools
        # Query notebook
        self.query_notebook = ctk.CTkTabview(self.query_frame, fg_color=theme_manager.get_color("editor.background"), 
                                       segmented_button_fg_color=theme_manager.get_color("sidebar.background"),
                                       segmented_button_selected_color=theme_manager.get_color("buttons.primary_bg"),
                                       segmented_button_selected_hover_color=theme_manager.get_color("buttons.primary_hover"),
                                       segmented_button_unselected_color=theme_manager.get_color("sidebar.background"),
                                       segmented_button_unselected_hover_color=theme_manager.get_color("sidebar.header"),
                                       text_color=theme_manager.get_color("text.primary"),
                                       text_color_disabled="#3E2723")
        self.query_notebook.pack(fill="both", expand=True, padx=0, pady=0)
        
        # Query tool tab
        self.query_notebook.add("Query Tool")
        self.query_panel = QueryPanel(
            self.query_notebook.tab("Query Tool"), 
            self.execute_query_callback,
            self.ai_generate_callback,
            self.display_results_callback,
            None  # Schema browser will be set after creation
        )
        self.query_panel.pack(fill="both", expand=True)
        
        # Now create schema browser with query callback and connection callbacks
        self.schema_browser = SchemaBrowser(
            self.left_frame, 
            self.on_table_select,
            self.on_saved_query_select,
            self.ai_assistant,
            self.show_connection_dialog,  # on_connect
            self.disconnect_database,     # on_disconnect
            self.db_connection            # db_connection
        )
        self.schema_browser.pack(fill="both", expand=True, padx=0, pady=0)
        
        # Link query panel to schema browser
        self.query_panel.set_schema_browser(self.schema_browser)
        
        # PSQL terminal tab
        self.query_notebook.add("PSQL Terminal")
        self.psql_terminal = PSQLTerminal(self.query_notebook.tab("PSQL Terminal"))
        self.psql_terminal.pack(fill="both", expand=True)
        
        # Results panel
        self.results_frame = ctk.CTkFrame(self.right_paned, height=400, corner_radius=8)
        self.right_paned.add(self.results_frame, height=400, minsize=200)
        
        # Create results display
        self.create_results_display()
    
    def create_results_display(self):
        """Create the results display area"""
        # Configure grid
        self.results_frame.grid_columnconfigure(0, weight=1)
        self.results_frame.grid_rowconfigure(1, weight=1)
        
        # Results header
        results_header = ctk.CTkFrame(self.results_frame, height=45, fg_color=theme_manager.get_color("card"), corner_radius=8)
        results_header.grid(row=0, column=0, sticky="ew", padx=0, pady=0)
        results_header.grid_columnconfigure(0, weight=1)
        
        self.results_label = ctk.CTkLabel(
            results_header, 
            text="Query results will appear here", 
            font=ctk.CTkFont(size=15, weight="bold"),
            text_color=theme_manager.get_color("foreground")
        )
        self.results_label.grid(row=0, column=0, sticky="w", padx=15, pady=8)
        
        # Export buttons frame
        export_frame = ctk.CTkFrame(results_header, fg_color="transparent")
        export_frame.grid(row=0, column=1, padx=15, pady=8)
        
        # CSV Export button
        self.export_csv_btn = ctk.CTkButton(
            export_frame, 
            text="📄 CSV", 
            command=self.export_csv,
            width=80,
            height=32,
            state="disabled",
            fg_color=theme_manager.get_color("buttons.primary_bg"),
            hover_color=theme_manager.get_color("buttons.primary_hover"),
            corner_radius=6
        )
        self.export_csv_btn.grid(row=0, column=0, padx=(0, 5))
        
        # Excel Export button
        self.export_excel_btn = ctk.CTkButton(
            export_frame, 
            text="📊 Excel", 
            command=self.export_excel,
            width=80,
            height=32,
            state="disabled",
            fg_color=theme_manager.get_color("buttons.primary_bg"),
            hover_color=theme_manager.get_color("buttons.primary_hover"),
            corner_radius=6
        )
        self.export_excel_btn.grid(row=0, column=1, padx=(5, 0))
        
        # Results table frame
        table_frame = ctk.CTkFrame(self.results_frame, corner_radius=8)
        table_frame.grid(row=1, column=0, sticky="nsew", padx=0, pady=0)
        table_frame.grid_columnconfigure(0, weight=1)
        table_frame.grid_rowconfigure(0, weight=1)
        
        # Create treeview for results with tree column for row numbers
        self.results_tree = ttk.Treeview(
            table_frame, 
            show="tree headings",
            selectmode="extended"
        )
        
        # Scrollbars for results
        results_v_scroll = ttk.Scrollbar(table_frame, orient="vertical", command=self.results_tree.yview)
        results_h_scroll = ttk.Scrollbar(table_frame, orient="horizontal", command=self.results_tree.xview)
        self.results_tree.configure(yscrollcommand=results_v_scroll.set, xscrollcommand=results_h_scroll.set)
        
        # Pack results table and scrollbars
        self.results_tree.grid(row=0, column=0, sticky="nsew", padx=(1, 0), pady=(1, 0))
        results_v_scroll.grid(row=0, column=1, sticky="ns", pady=(1, 0))
        results_h_scroll.grid(row=1, column=0, sticky="ew", padx=(1, 0))
        
        # Bind cell selection events
        self.results_tree.bind("<Button-1>", self.on_results_cell_click)
        self.results_tree.bind("<Double-1>", self.on_results_cell_double_click)
        self.results_tree.bind("<Button-3>", self.on_results_right_click)
        
        # Bind click on empty area to close popover
        self.results_tree.bind("<Button-1>", self.on_results_area_click, add=True)
        
        # Cell selection state
        self.selected_cell_row = None
        self.selected_cell_column = None
        self.selected_cell_value = None
        self.cell_popover = None  # For tracking the popover window
        
        # Configure results styling
        self.configure_results_style()
        
        # Current results data
        self.current_results = []
        self.current_columns = []
    
    def create_status_bar(self):
        """Create the status bar"""
        self.status_frame = ctk.CTkFrame(self, height=30, fg_color=theme_manager.get_color("card"))
        self.status_frame.grid(row=1, column=0, sticky="ew", padx=5, pady=(0, 5))
        self.status_frame.grid_columnconfigure(0, weight=1)
        
        self.status_label = ctk.CTkLabel(
            self.status_frame, 
            text="Ready", 
            font=ctk.CTkFont(size=11),
            text_color=theme_manager.get_color("foreground")
        )
        self.status_label.grid(row=0, column=0, padx=10, pady=5, sticky="w")
    
    def configure_results_style(self):
        """Configure results table styling"""
        style = ttk.Style()
        
        # Configure treeview colors using theme
        style.configure("Treeview", 
                        background=theme_manager.get_color("background"),
                        foreground=theme_manager.get_color("foreground"),
                        fieldbackground=theme_manager.get_color("background"),
                        borderwidth=1,
                        font=("Consolas", 11),
                        rowheight=25)
        style.configure("Treeview.Heading",
                        background=theme_manager.get_color("card"),
                        foreground=theme_manager.get_color("foreground"),
                        borderwidth=1,
                        relief="raised",
                        font=("Consolas", 11, "bold"))
        style.map("Treeview",
                  background=[('selected', theme_manager.get_color("primary"))],
                  foreground=[('selected', theme_manager.get_color("primaryForeground"))])
        
        # Configure scrollbars using theme
        style.configure("Vertical.TScrollbar", 
                        background=theme_manager.get_color("card"), 
                        troughcolor=theme_manager.get_color("background"), 
                        borderwidth=1,
                        arrowcolor=theme_manager.get_color("foreground"))
        style.configure("Horizontal.TScrollbar", 
                        background=theme_manager.get_color("card"), 
                        troughcolor=theme_manager.get_color("background"), 
                        borderwidth=1,
                        arrowcolor=theme_manager.get_color("foreground"))
    
    def on_results_area_click(self, event):
        """Handle clicks on results area (to close popover when clicking outside cells)"""
        region = self.results_tree.identify_region(event.x, event.y)
        if region not in ["cell"]:
            # Clicked on empty area, close popover
            self.close_cell_popover()

    def on_results_cell_click(self, event):
        """Handle single click on results table cell"""
        # Identify which cell was clicked
        region = self.results_tree.identify_region(event.x, event.y)
        
        if region == "cell":
            # Get the row
            row_id = self.results_tree.identify_row(event.y)
            if not row_id:
                return
            
            # Get the column
            column_id = self.results_tree.identify_column(event.x)
            if not column_id:
                return
            
            # Convert column_id to column name
            if column_id == "#0":
                # Clicked on row number column
                return
            
            # Get column index (columns are #1, #2, etc.)
            col_index = int(column_id.replace("#", "")) - 1
            if col_index < 0 or col_index >= len(self.current_columns):
                return
            
            column_name = self.current_columns[col_index]
            
            # Get the row data
            row_index = int(self.results_tree.item(row_id, "text")) - 1
            if row_index < 0 or row_index >= len(self.current_results):
                return
            
            # Get cell value
            cell_value = self.current_results[row_index].get(column_name, "")
            
            # Store selected cell info
            self.selected_cell_row = row_index
            self.selected_cell_column = column_name
            self.selected_cell_value = cell_value
            
            # Update status bar to show selected cell info (single click just selects)
            value_preview = str(cell_value)
            if len(value_preview) > 50:
                value_preview = value_preview[:47] + "..."
            
            self.status_label.configure(
                text=f"Selected: Row {row_index + 1}, {column_name} = {value_preview}"
            )
            
            # Close any existing popover
            self.close_cell_popover()
    
    def on_results_cell_click_select_only(self, event):
        """Handle cell click for selection only (used by double-click)"""
        # Identify which cell was clicked
        region = self.results_tree.identify_region(event.x, event.y)
        
        if region == "cell":
            # Get the row
            row_id = self.results_tree.identify_row(event.y)
            if not row_id:
                return
            
            # Get the column
            column_id = self.results_tree.identify_column(event.x)
            if not column_id:
                return
            
            # Convert column_id to column name
            if column_id == "#0":
                # Clicked on row number column
                return
            
            # Get column index (columns are #1, #2, etc.)
            col_index = int(column_id.replace("#", "")) - 1
            if col_index < 0 or col_index >= len(self.current_columns):
                return
            
            column_name = self.current_columns[col_index]
            
            # Get the row data
            row_index = int(self.results_tree.item(row_id, "text")) - 1
            if row_index < 0 or row_index >= len(self.current_results):
                return
            
            # Get cell value
            cell_value = self.current_results[row_index].get(column_name, "")
            
            # Store selected cell info
            self.selected_cell_row = row_index
            self.selected_cell_column = column_name
            self.selected_cell_value = cell_value

    def on_results_cell_double_click(self, event):
        """Handle double-click on results table cell to show popover"""
        # First ensure cell is selected by calling the selection handler
        self.on_results_cell_click_select_only(event)
        
        if self.selected_cell_value is not None:
            # Show the cell value popover next to the clicked cell
            self.show_cell_popover(event)
    
    def on_results_right_click(self, event):
        """Handle right-click on results table cell - show context menu"""
        # First, select the cell
        self.on_results_cell_click(event)
        
        if self.selected_cell_value is None:
            return
        
        # Create context menu
        context_menu = tk.Menu(self, tearoff=0)
        context_menu.configure(
            background=theme_manager.get_color("background.main"),
            foreground=theme_manager.get_color("text.primary"),
            activebackground=theme_manager.get_color("buttons.primary_bg"),
            activeforeground=theme_manager.get_color("buttons.primary_text"),
            font=("Segoe UI", 10)
        )
        
        context_menu.add_command(
            label=f"📋 Copy Cell ({self.selected_cell_column})",
            command=self.copy_selected_cell
        )
        
        context_menu.add_command(
            label="📄 Copy Row",
            command=self.copy_selected_row
        )
        
        context_menu.add_separator()
        
        context_menu.add_command(
            label="📊 View Full Value",
            command=self.view_full_cell_value
        )
        
        # Show menu at cursor position
        try:
            context_menu.tk_popup(event.x_root, event.y_root)
        finally:
            context_menu.grab_release()
    
    def copy_selected_cell(self):
        """Copy the selected cell value to clipboard"""
        if self.selected_cell_value is not None:
            self.clipboard_clear()
            self.clipboard_append(str(self.selected_cell_value))
            self.status_label.configure(text=f"✓ Copied: {self.selected_cell_column}")
            self.after(2000, lambda: self.status_label.configure(text="Ready"))
    
    def copy_selected_row(self):
        """Copy the entire selected row to clipboard"""
        if self.selected_cell_row is not None and self.selected_cell_row < len(self.current_results):
            row_data = self.current_results[self.selected_cell_row]
            
            # Format as tab-separated values
            row_values = [str(row_data.get(col, "")) for col in self.current_columns]
            row_text = "\t".join(row_values)
            
            self.clipboard_clear()
            self.clipboard_append(row_text)
            self.status_label.configure(text=f"✓ Copied row {self.selected_cell_row + 1}")
            self.after(2000, lambda: self.status_label.configure(text="Ready"))
    
    def show_cell_popover(self, event):
        """Show a popover next to the clicked cell"""
        if self.selected_cell_value is None:
            return
        
        # Close any existing popover
        if hasattr(self, 'cell_popover') and self.cell_popover:
            self.cell_popover.destroy()
        
        # Get the bounding box of the clicked cell
        row_id = self.results_tree.identify_row(event.y)
        column_id = self.results_tree.identify_column(event.x)
        
        if not row_id or not column_id:
            return
        
        # Get cell position relative to the treeview
        cell_bbox = self.results_tree.bbox(row_id, column_id)
        if not cell_bbox:
            return
        
        cell_x, cell_y, cell_width, cell_height = cell_bbox
        
        # Convert to screen coordinates
        tree_x = self.results_tree.winfo_rootx()
        tree_y = self.results_tree.winfo_rooty()
        
        # Position popover to the right of the cell
        popover_x = tree_x + cell_x + cell_width + 10
        popover_y = tree_y + cell_y
        
        # Create popover window
        self.cell_popover = tk.Toplevel(self)
        self.cell_popover.wm_overrideredirect(True)  # Remove window decorations
        self.cell_popover.configure(bg=theme_manager.get_color("background.main"))
        
        # Set initial size
        popover_width = 350
        popover_height = 250  # Increased height to ensure buttons are visible
        
        # Check screen bounds and adjust position
        screen_width = self.winfo_screenwidth()
        screen_height = self.winfo_screenheight()
        
        # Adjust horizontal position if it would go off screen
        if popover_x + popover_width > screen_width:
            popover_x = tree_x + cell_x - popover_width - 10  # Position to the left instead
        
        # Adjust vertical position if it would go off screen
        if popover_y + popover_height > screen_height:
            popover_y = tree_y + cell_y + cell_height - popover_height
        
        # Ensure popover doesn't go above screen or too far left
        if popover_y < 0:
            popover_y = 10
        if popover_x < 0:
            popover_x = 10
        
        self.cell_popover.geometry(f"{popover_width}x{popover_height}+{popover_x}+{popover_y}")
        
        # Create main frame with border
        main_frame = ctk.CTkFrame(
            self.cell_popover,
            fg_color=theme_manager.get_color("card"),
            border_width=2,
            border_color=theme_manager.get_color("accent.main")
        )
        main_frame.pack(fill="both", expand=True, padx=2, pady=2)
        
        # Header with column name and close button
        header_frame = ctk.CTkFrame(main_frame, fg_color=theme_manager.get_color("buttons.primary_bg"))
        header_frame.pack(fill="x", padx=5, pady=5)
        
        header_label = ctk.CTkLabel(
            header_frame,
            text=f"📊 {self.selected_cell_column} (Row {self.selected_cell_row + 1})",
            font=ctk.CTkFont(size=11, weight="bold"),
            text_color=theme_manager.get_color("buttons.primary_text")
        )
        header_label.pack(side="left", padx=8, pady=4)
        
        close_btn = ctk.CTkButton(
            header_frame,
            text="✕",
            width=20,
            height=20,
            command=self.close_cell_popover,
            fg_color="transparent",
            hover_color=theme_manager.get_color("button.hoverBackground"),
            text_color=theme_manager.get_color("buttons.primary_text")
        )
        close_btn.pack(side="right", padx=4, pady=2)
        
        # Content area with scrollable text
        content_frame = ctk.CTkFrame(main_frame, fg_color="transparent")
        content_frame.pack(fill="both", expand=True, padx=5, pady=(0, 5))
        
        # Text widget for content
        text_widget = tk.Text(
            content_frame,
            wrap="word",
            bg=theme_manager.get_color("editor.background"),
            fg=theme_manager.get_color("text.primary"),
            font=("Consolas", 10),
            padx=8,
            pady=8,
            height=6,  # Reduced height to make room for buttons
            relief="flat",
            borderwidth=0
        )
        text_widget.pack(side="left", fill="both", expand=True)
        
        # Scrollbar
        scrollbar = ttk.Scrollbar(content_frame, orient="vertical", command=text_widget.yview)
        scrollbar.pack(side="right", fill="y")
        text_widget.configure(yscrollcommand=scrollbar.set)
        
        # Insert the cell value
        text_widget.insert("1.0", str(self.selected_cell_value))
        text_widget.configure(state="disabled")
        
        # Bottom buttons
        button_frame = ctk.CTkFrame(main_frame, fg_color=theme_manager.get_color("sidebar.background"), height=40)
        button_frame.pack(fill="x", padx=5, pady=5)
        button_frame.pack_propagate(False)  # Maintain fixed height
        
        copy_btn = ctk.CTkButton(
            button_frame,
            text="📋 Copy",
            command=self.copy_cell_value_and_close,
            fg_color=theme_manager.get_color("buttons.primary_bg"),
            hover_color=theme_manager.get_color("buttons.primary_hover"),
            text_color=theme_manager.get_color("buttons.primary_text"),
            width=90,
            height=30,
            font=ctk.CTkFont(size=11)
        )
        copy_btn.pack(side="left", padx=5, pady=5)
        
        view_full_btn = ctk.CTkButton(
            button_frame,
            text="🔍 View Full",
            command=self.open_full_cell_dialog,
            fg_color=theme_manager.get_color("button.background"),
            hover_color=theme_manager.get_color("button.hoverBackground"),
            text_color=theme_manager.get_color("text.primary"),
            width=90,
            height=30,
            font=ctk.CTkFont(size=11)
        )
        view_full_btn.pack(side="left", padx=5, pady=5)
        
        # Make popover closable by clicking outside
        self.cell_popover.bind("<FocusOut>", lambda e: self.close_cell_popover())
        self.cell_popover.focus_set()
        
        # Close on Escape key
        self.cell_popover.bind("<Escape>", lambda e: self.close_cell_popover())
    
    def close_cell_popover(self):
        """Close the cell popover"""
        if hasattr(self, 'cell_popover') and self.cell_popover:
            self.cell_popover.destroy()
            self.cell_popover = None
    
    def copy_cell_value_and_close(self):
        """Copy cell value to clipboard and close popover"""
        if self.selected_cell_value is not None:
            self.clipboard_clear()
            self.clipboard_append(str(self.selected_cell_value))
            self.status_label.configure(text=f"✓ Copied: {self.selected_cell_column}")
            self.after(2000, lambda: self.status_label.configure(text="Ready"))
        self.close_cell_popover()
    
    def open_full_cell_dialog(self):
        """Open the full cell value dialog and close popover"""
        self.close_cell_popover()
        self.view_full_cell_value()

    def view_full_cell_value(self):
        """Show the full cell value in a dialog"""
        if self.selected_cell_value is None:
            return
        
        # Create a dialog to show full value
        dialog = tk.Toplevel(self)
        dialog.title(f"Cell Value: {self.selected_cell_column}")
        dialog.geometry("600x400")
        dialog.configure(bg=theme_manager.get_color("background.main"))
        
        # Make it modal
        dialog.transient(self)
        dialog.grab_set()
        
        # Title label
        title_frame = ctk.CTkFrame(dialog, fg_color=theme_manager.get_color("buttons.primary_bg"))
        title_frame.pack(fill="x", padx=0, pady=0)
        
        title_label = ctk.CTkLabel(
            title_frame,
            text=f"📊 Row {self.selected_cell_row + 1} • {self.selected_cell_column}",
            font=ctk.CTkFont(size=12, weight="bold"),
            text_color=theme_manager.get_color("buttons.primary_text")
        )
        title_label.pack(pady=8, padx=12)
        
        # Text widget to show value
        text_frame = tk.Frame(dialog, bg=theme_manager.get_color("editor.background"))
        text_frame.pack(fill="both", expand=True, padx=12, pady=12)
        
        text_widget = tk.Text(
            text_frame,
            wrap="word",
            bg=theme_manager.get_color("editor.background"),
            fg=theme_manager.get_color("text.primary"),
            font=("Consolas", 11),
            padx=10,
            pady=10
        )
        text_widget.pack(side="left", fill="both", expand=True)
        
        # Scrollbar
        scrollbar = ttk.Scrollbar(text_frame, orient="vertical", command=text_widget.yview)
        scrollbar.pack(side="right", fill="y")
        text_widget.configure(yscrollcommand=scrollbar.set)
        
        # Insert value
        text_widget.insert("1.0", str(self.selected_cell_value))
        text_widget.configure(state="disabled")
        
        # Button frame
        button_frame = ctk.CTkFrame(dialog, fg_color="transparent")
        button_frame.pack(fill="x", padx=12, pady=(0, 12))
        
        copy_btn = ctk.CTkButton(
            button_frame,
            text="📋 Copy",
            command=lambda: [self.clipboard_clear(), 
                           self.clipboard_append(str(self.selected_cell_value)),
                           self.status_label.configure(text="✓ Copied to clipboard")],
            fg_color=theme_manager.get_color("button.background"),
            hover_color=theme_manager.get_color("button.hoverBackground"),
            width=100
        )
        copy_btn.pack(side="left", padx=5)
        
        close_btn = ctk.CTkButton(
            button_frame,
            text="Close",
            command=dialog.destroy,
            fg_color=theme_manager.get_color("button.hoverBackground"),
            hover_color=theme_manager.get_color("sideBarSectionHeader.background"),
            width=100
        )
        close_btn.pack(side="left", padx=5)
    
    def display_results_callback(self, results, columns):
        """Callback to display results from query panel"""
        self.display_results(results, columns)
    
    def display_results(self, results, columns):
        """Display query results in the main results area"""
        # Clear existing results
        self.clear_results()
        
        if not results or not columns:
            self.results_label.configure(text="Query results will appear here")
            return
        
        # Store current results
        self.current_results = results
        self.current_columns = columns
        
        # Configure tree column for row numbers
        self.results_tree.column("#0", width=50, minwidth=50, anchor="center", stretch=False)
        self.results_tree.heading("#0", text="#", anchor="center")
        
        # Configure columns
        self.results_tree["columns"] = columns
        
        # Set column headings and widths
        for col in columns:
            self.results_tree.heading(col, text=col, anchor="w")
            # Calculate column width based on content
            header_width = len(col) * 12
            max_width = max(120, header_width)
            
            if results:
                # Sample first few rows to estimate width
                sample_rows = results[:min(20, len(results))]
                for row in sample_rows:
                    if col in row and row[col] is not None:
                        content_length = len(str(row[col]))
                        if content_length <= 10:
                            content_width = content_length * 12
                        elif content_length <= 50:
                            content_width = content_length * 10
                        else:
                            content_width = 300
                        max_width = max(max_width, content_width)
            
            # Cap maximum and minimum widths
            max_width = min(max(max_width, 80), 400)
            self.results_tree.column(col, width=max_width, anchor="w", minwidth=80)
        
        # Insert data
        for i, row in enumerate(results):
            values = []
            for col in columns:
                value = row.get(col, "")
                # Handle None values
                if value is None:
                    value = "[NULL]"
                # Handle different data types
                elif isinstance(value, (list, dict)):
                    value = str(value)
                    if len(value) > 200:
                        value = value[:197] + "..."
                # Truncate very long text values for display
                elif isinstance(value, str):
                    if len(value) > 200:
                        value = value[:197] + "..."
                    # Replace newlines and tabs for better display
                    value = value.replace('\n', ' ').replace('\t', ' ').replace('\r', ' ')
                # Handle other types
                else:
                    value = str(value)
                    if len(value) > 200:
                        value = value[:197] + "..."
                
                values.append(value)
            
            # Add alternating row colors by using tags
            tag = "odd" if i % 2 == 1 else "even"
            # Insert with row number in the tree column
            self.results_tree.insert("", "end", text=str(i + 1), values=values, tags=(tag,))
        
        # Configure row tags for better readability using theme colors
        self.results_tree.tag_configure("odd", background=theme_manager.get_color("table.background"))
        self.results_tree.tag_configure("even", background=theme_manager.get_color("background.secondary"))
        
        # Update results label
        self.results_label.configure(text=f"Results ({len(results)} rows)")
        
        # Enable export buttons
        self.export_csv_btn.configure(state="normal" if results else "disabled")
        self.export_excel_btn.configure(state="normal" if results else "disabled")
    
    def clear_results(self):
        """Clear the results table"""
        for item in self.results_tree.get_children():
            self.results_tree.delete(item)
        
        self.results_tree["columns"] = ()
        self.results_label.configure(text="Query results will appear here")
        self.export_csv_btn.configure(state="disabled")
        self.export_excel_btn.configure(state="disabled")
        
        # Clear stored data
        self.current_results = []
        self.current_columns = []
    
    def export_csv(self):
        """Export query results to CSV"""
        if not self.current_results or not self.current_columns:
            messagebox.showwarning("No Results", "No query results to export")
            return
        
        from tkinter import filedialog
        
        # Ask for save location
        filename = filedialog.asksaveasfilename(
            defaultextension=".csv",
            filetypes=[("CSV files", "*.csv"), ("All files", "*.*")],
            title="Export Results"
        )
        
        if not filename:
            return
        
        try:
            import csv
            
            with open(filename, 'w', newline='', encoding='utf-8') as csvfile:
                writer = csv.DictWriter(csvfile, fieldnames=self.current_columns)
                writer.writeheader()
                
                for row in self.current_results:
                    # Handle None values and convert all to strings
                    clean_row = {}
                    for col in self.current_columns:
                        value = row.get(col)
                        clean_row[col] = "" if value is None else str(value)
                    writer.writerow(clean_row)
            
            messagebox.showinfo("Export Complete", f"Results exported to:\n{filename}")
            
        except Exception as e:
            messagebox.showerror("Export Error", f"Failed to export results:\n{e}")
    
    def export_excel(self):
        """Export query results to Excel"""
        if not self.current_results or not self.current_columns:
            messagebox.showwarning("No Results", "No query results to export")
            return
        
        # Ask for save location
        filename = filedialog.asksaveasfilename(
            defaultextension=".xlsx",
            filetypes=[("Excel files", "*.xlsx"), ("All files", "*.*")],
            title="Export Results to Excel"
        )
        
        if not filename:
            return
        
        try:
            # Convert results to DataFrame
            data = []
            for row in self.current_results:
                clean_row = {}
                for col in self.current_columns:
                    value = row.get(col)
                    clean_row[col] = value  # Keep original types for Excel
                data.append(clean_row)
            
            df = pd.DataFrame(data, columns=self.current_columns)
            
            # Create Excel writer with formatting
            with pd.ExcelWriter(filename, engine='openpyxl') as writer:
                # Write data to Excel
                df.to_excel(writer, sheet_name='Query Results', index=False)
                
                # Get the workbook and worksheet
                workbook = writer.book
                worksheet = writer.sheets['Query Results']
                
                # Auto-adjust column widths
                for column in worksheet.columns:
                    max_length = 0
                    column_letter = column[0].column_letter
                    
                    for cell in column:
                        try:
                            if len(str(cell.value)) > max_length:
                                max_length = len(str(cell.value))
                        except:
                            pass
                    
                    adjusted_width = min(max_length + 2, 50)  # Max width of 50
                    worksheet.column_dimensions[column_letter].width = adjusted_width
                
                # Style the header row
                from openpyxl.styles import Font, PatternFill, Alignment
                
                header_font = Font(bold=True, color="FFFFFF")
                header_fill = PatternFill(start_color="0078D4", end_color="0078D4", fill_type="solid")
                center_alignment = Alignment(horizontal="center", vertical="center")
                
                for cell in worksheet[1]:  # Header row
                    cell.font = header_font
                    cell.fill = header_fill
                    cell.alignment = center_alignment
            
            messagebox.showinfo("Export Complete", f"Results exported to Excel:\n{filename}")
            
        except ImportError:
            messagebox.showerror("Export Error", "Excel export requires 'openpyxl' package.\nPlease install it with: pip install openpyxl")
        except Exception as e:
            messagebox.showerror("Export Error", f"Failed to export results to Excel:\n{e}")
    
    def init_ai_assistant(self):
        """Initialize the AI assistant"""
        try:
            self.ai_assistant = NeuronDBAI()
            self.logger.info("AI assistant initialized successfully")
        except Exception as e:
            self.logger.error(f"Failed to initialize AI assistant: {e}")
            messagebox.showwarning(
                "AI Assistant", 
                f"Failed to initialize AI assistant:\n{e}\n\n"
                "Please check your GOOGLE_API_KEY in .env file"
            )
    
    def show_connection_dialog(self):
        """Show connection dialog"""
        dialog = ConnectionDialog(self, self.connection_manager)
        self.wait_window(dialog)
        
        # Check if a connection was made
        if hasattr(dialog, 'selected_connection') and dialog.selected_connection:
            self.connect_to_database(dialog.selected_connection)
    
    def connect_to_database(self, connection_config: Dict[str, Any]):
        """Connect to database with given configuration"""
        def _connect_thread():
            """Thread function for database connection"""
            try:
                self.logger.info(f"[UI] Starting database connection process...")
                self.logger.info(f"[UI] Config: {connection_config['database']}@{connection_config['host']}:{connection_config['port']}")
                
                # Attempt connection
                self.logger.info("[UI] Calling db_connection.connect()...")
                self.db_connection.connect(
                    host=connection_config['host'],
                    port=connection_config['port'],
                    database=connection_config['database'],
                    username=connection_config['username'],
                    password=connection_config['password']
                )
                self.logger.info("[UI] Connection successful, updating UI...")
                
                # Update UI in main thread
                self.after(0, self._on_connection_success, connection_config)
                
            except Exception as e:
                self.logger.error(f"[UI] ❌ Connection failed: {type(e).__name__}: {e}")
                # Update UI in main thread
                self.after(0, self._on_connection_error, str(e))
        
        # Update status and start connection in background thread
        self.update_status("Connecting to database...")
        thread = threading.Thread(target=_connect_thread, daemon=True)
        thread.start()
    
    def _on_connection_success(self, connection_config: Dict[str, Any]):
        """Called in main thread after successful connection"""
        try:
            # Update UI - update schema browser connection state with DB info
            db_info = f"{connection_config['database']}@{connection_config['host']}:{connection_config['port']}"
            self.schema_browser.set_connected(True, db_info)
            
            # Load schema in background
            self.logger.info("[UI] Refreshing schema...")
            self.refresh_schema()
            
            # Update PSQL terminal
            self.logger.info("[UI] Updating PSQL terminal...")
            self.psql_terminal.set_connection(self.db_connection)
            
            self.update_status("Connected successfully")
            self.logger.info(f"[UI] ✅ Connected to database: {connection_config['database']}")
            
        except Exception as e:
            self.logger.error(f"[UI] Error in post-connection setup: {e}")
            self.update_status("Connection succeeded but setup failed")
    
    def _on_connection_error(self, error_message: str):
        """Called in main thread after connection failure"""
        messagebox.showerror("Connection Error", f"Failed to connect to database:\n{error_message}")
        self.update_status("Connection failed")
    
    def disconnect_database(self):
        """Disconnect from current database"""
        try:
            self.logger.info("[UI] Disconnecting from database...")
            
            if self.db_connection and self.db_connection.is_connected():
                self.db_connection.disconnect()
                self.logger.info("[UI] Database connection closed")
            
            # Update UI - update schema browser connection state
            self.schema_browser.set_connected(False)
            
            # Clear schema
            self.current_schema = {}
            self.schema_browser.clear_schema()
            
            # Clear AI schema context
            if self.ai_assistant:
                self.ai_assistant.set_database_schema({})
            
            # Update PSQL terminal
            self.psql_terminal.clear_connection()
            
            self.update_status("Disconnected")
            self.logger.info("[UI] ✅ Disconnected from database successfully")
            
        except Exception as e:
            self.logger.error(f"[UI] ❌ Disconnect failed: {type(e).__name__}: {e}")
            messagebox.showerror("Disconnect Error", f"Failed to disconnect:\n{e}")
    
    def refresh_schema(self):
        """Refresh database schema"""
        if not self.db_connection.is_connected():
            messagebox.showwarning("Not Connected", "Please connect to a database first")
            return
        
        def _refresh_schema_thread():
            """Thread function for schema refresh"""
            try:
                self.logger.info("[UI] Starting schema refresh...")
                
                # Get schema from database (this can be slow)
                self.logger.info("[UI] Calling get_database_schema()...")
                schema = self.db_connection.get_database_schema()
                
                self.logger.info(f"[UI] Schema loaded: {len(schema.get('tables', {}))} tables")
                
                # Update UI in main thread
                self.after(0, self._on_schema_loaded, schema)
                
            except Exception as e:
                self.logger.error(f"[UI] Schema refresh failed: {type(e).__name__}: {e}")
                # Update UI in main thread
                self.after(0, self._on_schema_error, str(e))
        
        # Update status and start schema fetch in background thread
        self.update_status("Loading database schema...")
        thread = threading.Thread(target=_refresh_schema_thread, daemon=True)
        thread.start()
    
    def _on_schema_loaded(self, schema: Dict[str, Any]):
        """Called in main thread after schema is loaded"""
        try:
            self.current_schema = schema
            
            # Update schema browser
            self.logger.info("[UI] Updating schema browser...")
            self.schema_browser.update_schema(self.current_schema)
            
            # Update AI assistant with schema context
            self.logger.info("[UI] Updating AI assistant schema...")
            if self.ai_assistant:
                self.ai_assistant.set_database_schema(self.current_schema)
            
            table_count = len(self.current_schema.get('tables', {}))
            view_count = len(self.current_schema.get('views', {}))
            self.update_status(f"Schema loaded: {table_count} tables, {view_count} views")
            self.logger.info(f"[UI] ✅ Schema refresh complete")
            
        except Exception as e:
            self.logger.error(f"[UI] Error updating UI with schema: {e}")
            self.update_status("Schema loaded but UI update failed")
    
    def _on_schema_error(self, error_message: str):
        """Called in main thread after schema loading failure"""
        messagebox.showerror("Schema Error", f"Failed to load database schema:\n{error_message}")
        self.update_status("Schema load failed")
    
    def execute_current_query(self):
        """Execute the current query from query panel"""
        if hasattr(self.query_panel, 'execute_query'):
            self.query_panel.execute_query()
    
    def clear_query(self):
        """Clear the current query"""
        if hasattr(self.query_panel, 'clear_query'):
            self.query_panel.clear_query()
    
    def format_current_query(self):
        """Format the current query"""
        if hasattr(self.query_panel, 'format_query'):
            self.query_panel.format_query()
    
    def show_ai_dialog(self):
        """Show AI query generation dialog"""
        if not self.ai_assistant:
            messagebox.showwarning("AI Not Available", "AI assistant is not configured")
            return
        
        # This would open an AI chat dialog - simplified for now
        user_input = tk.simpledialog.askstring(
            "AI Query Generator",
            "Describe what you want to query:"
        )
        
        if user_input and hasattr(self.query_panel, 'generate_with_ai'):
            self.query_panel.generate_with_ai(user_input)
    
    def explain_current_query(self):
        """Explain the current query using AI"""
        if hasattr(self.query_panel, 'explain_query'):
            self.query_panel.explain_query()
    
    def clear_ai_history(self):
        """Clear AI conversation history"""
        if self.ai_assistant:
            self.ai_assistant.clear_conversation_history()
            messagebox.showinfo("AI History", "Conversation history cleared")
    
    def show_about(self):
        """Show about dialog"""
        about_text = """NeuronDB - AI-Powered PostgreSQL Client

Version: 1.0.0

A modern desktop application that brings AI assistance to PostgreSQL database management.

Features:
• AI-powered query generation with Google Gemini
• Visual schema browser
• Advanced query editor with saved queries
• Built-in PSQL terminal
• Connection management
• Excel export support

Built with Python, CustomTkinter, and LangChain.
"""
        messagebox.showinfo("About NeuronDB", about_text)
    
    def execute_query_callback(self, query: str):
        """Callback for executing queries from query panel"""
        if not self.db_connection.is_connected():
            return None, "Not connected to database"
        
        try:
            results, columns = self.db_connection.execute_query(query)
            return results, columns
        except Exception as e:
            return None, str(e)
    
    def execute_query(self):
        """Execute the current query in the query panel"""
        if hasattr(self.query_panel, 'execute_query'):
            self.query_panel.execute_query()
    
    def ai_generate_callback(self, user_input: str):
        """Callback for AI query generation"""
        if not self.ai_assistant:
            return None, "AI assistant not available"
        
        try:
            result = self.ai_assistant.generate_sql_query(user_input)
            if result['success']:
                return result['query'], None
            else:
                return None, result['error']
        except Exception as e:
            return None, str(e)
    
    def on_table_select(self, table_name: str):
        """Handle table selection from schema browser"""
        if hasattr(self.query_panel, 'insert_table_name'):
            self.query_panel.insert_table_name(table_name)
    
    def on_saved_query_select(self, query_text: str):
        """Handle saved query selection from schema browser"""
        if hasattr(self.query_panel, 'append_query'):
            self.query_panel.append_query(query_text)
            self.update_status("Saved query appended to editor")
    
    def update_status(self, message: str):
        """Update status bar message"""
        self.status_label.configure(text=message)
        self.update_idletasks()
    
    def on_closing(self):
        """Handle application closing"""
        try:
            self.logger.info("Application closing initiated...")
            
            # Disconnect from database if connected
            if self.db_connection and self.db_connection.is_connected():
                self.logger.info("Closing database connection...")
                self.update_status("Disconnecting from database...")
                self.db_connection.disconnect()
                self.logger.info("✅ Database connection closed successfully")
            
            # Clean up AI assistant if exists
            if self.ai_assistant:
                self.logger.info("Cleaning up AI assistant...")
                self.ai_assistant = None
            
            self.logger.info("NeuronDB application closed")
            self.destroy()
            
        except Exception as e:
            self.logger.error(f"❌ Error during shutdown: {type(e).__name__}: {e}")
            # Force close even if there's an error
            try:
                self.destroy()
            except:
                pass

if __name__ == "__main__":
    app = NeuronDBApp()
    app.mainloop()